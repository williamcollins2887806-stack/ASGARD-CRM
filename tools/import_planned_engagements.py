#!/usr/bin/env python3
"""
Импорт планируемого привлечения из двух Excel:
  1) Персонал на останов — ФИО + объект на останов
  2) График перевахтовки — planned_from / planned_to по календарю (11 = рабочий день)

Использование:
  python tools/import_planned_engagements.py --dry-run
  python tools/import_planned_engagements.py --apply
"""
from __future__ import annotations

import argparse
import io
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

import openpyxl
import paramiko

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_PERSONNEL = Path(r"C:\Users\Nikita-ASGARD\Downloads\Персонал на останов 2026г от 10.07.26.xlsx")
DEFAULT_SCHEDULE = Path(r"C:\Users\Nikita-ASGARD\Downloads\График перевахтовки от 06.07.2026.xlsx")
SSH_KEY = Path.home() / ".ssh" / "asgard_crm_deploy"
HOST = "92.242.61.184"
CREATED_BY = 3460  # Хосе

# Прямое сопоставление с works МЛСП Приразломная (site_id=1152)
WORK_IDS = {
    "химпромыв": 402,
    "восстанов": 354,
    "деаэратор": 403,
    "тк": 404,
    "каусорб": 405,
    "оголовок": 406,
    "блок котельн": 407,
}

SKIP_PROJECT_PATTERNS = (
    "по здоровью", "запрет", "семейн", "не едет", "не отвечает", "пуровск",
    "основной работе", "уход за", "???", "--------",
)

# Ключевые слова объекта Excel → ключ WORK_IDS
PROJECT_KEYWORDS = [
    ("химпромыв", "химпромыв"),
    ("буровые емкост", "тк"),
    ("тк емкост", "тк"),
    ("ёмкост", "тк"),
    ("емкост", "тк"),
    ("диаэратор", "деаэратор"),
    ("деаэратор", "деаэратор"),
    ("каусорб", "каусорб"),
    ("факельн", "оголовок"),
    ("оголовок", "оголовок"),
    ("блок котельн", "блок котельн"),
    (" бк", "блок котельн"),
    ("восстанов", "восстанов"),
    ("восст", "восстанов"),
]
def norm_phone(value) -> str:
    d = re.sub(r"\D", "", str(value or ""))
    if d.startswith("8") and len(d) == 11:
        d = "7" + d[1:]
    if len(d) == 10:
        d = "7" + d
    return d


def norm_fio(value: str) -> str:
    s = re.sub(r"\s+", " ", str(value or "").strip().lower())
    s = s.replace("ё", "е")
    return s


def norm_fio_key(value: str) -> str:
    """Сравнение по фамилии + инициалам."""
    s = norm_fio(value)
    s = re.sub(r"[^a-zа-я0-9.\s]", "", s)
    parts = s.split()
    if not parts:
        return ""
    fam = parts[0]
    initials = "".join(p[0] for p in parts[1:] if p)
    return f"{fam}|{initials}"


def parse_date_cell(value) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    s = str(value).strip()
    if not s or not re.search(r"\d{4}", s):
        return None
    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s[:10], fmt).date()
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(s[:19]).date()
    except ValueError:
        return None


def read_personnel_xlsx(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        if not row or not row[1]:
            continue
        fio = str(row[1]).strip()
        project = str(row[3] or "").strip()
        phone = norm_phone(row[4] if len(row) > 4 else "")
        arrival = parse_date_cell(row[7] if len(row) > 7 else None)
        rows.append({
            "fio": fio,
            "fio_key": norm_fio_key(fio),
            "phone_norm": phone,
            "project_raw": project,
            "planned_from_file": arrival,
            "row": i + 1,
        })
    wb.close()
    return rows


def work_day_ranges(dates: list[date]) -> list[tuple[date, date]]:
    if not dates:
        return []
    dates = sorted(set(dates))
    ranges: list[tuple[date, date]] = []
    start = prev = dates[0]
    for d in dates[1:]:
        if (d - prev).days == 1:
            prev = d
            continue
        ranges.append((start, prev))
        start = prev = d
    ranges.append((start, prev))
    return ranges


def read_schedule_xlsx(path: Path) -> dict[str, dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    hdr = list(next(ws.iter_rows(max_row=1, values_only=True)))
    date_cols: list[tuple[int, date]] = []
    for idx, val in enumerate(hdr):
        d = parse_date_cell(val)
        if d:
            date_cols.append((idx, d))

    by_fio: dict[str, dict] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[1]:
            continue
        fio = str(row[1]).strip()
        key = norm_fio_key(fio)
        work_dates: list[date] = []
        for idx, d in date_cols:
            if idx < len(row) and row[idx] in (11, "11", 11.0):
                work_dates.append(d)
        ranges = work_day_ranges(work_dates)
        # Берём последний диапазон с датой >= 2026-07-01 (останов), иначе последний вообще
        cutoff = date(2026, 7, 1)
        future = [r for r in ranges if r[0] >= cutoff]
        chosen = future[-1] if future else (ranges[-1] if ranges else None)
        by_fio[key] = {
            "fio": fio,
            "ranges": [(a.isoformat(), b.isoformat()) for a, b in ranges],
            "planned_from": chosen[0].isoformat() if chosen else None,
            "planned_to": chosen[1].isoformat() if chosen else None,
        }
    wb.close()
    return by_fio


def project_keyword(project_raw: str) -> str | None:
    low = norm_fio(project_raw)
    for needle, keyword in PROJECT_KEYWORDS:
        if needle in low:
            return keyword
    return None


def ssh_psql(sql: str) -> str:
    key = paramiko.Ed25519Key.from_private_key_file(str(SSH_KEY))
    client = paramiko.SSHClient()
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    client.connect(HOST, username="root", pkey=key, timeout=30)
    cmd = 'PGPASSWORD=123456789 psql -U asgard -d asgard_crm -t -A -F"|" -c "' + sql.replace('"', '\\"') + '"'
    _, stdout, stderr = client.exec_command(cmd, timeout=120)
    out = stdout.read().decode("utf-8", errors="replace").strip()
    err = stderr.read().decode("utf-8", errors="replace").strip()
    client.close()
    if err and "ERROR" in err.upper():
        raise RuntimeError(err)
    return out


def load_db_snapshot() -> tuple[list[dict], list[dict]]:
    emp_sql = """
      SELECT id, fio, phone FROM employees
      WHERE COALESCE(is_active, true) = true
      ORDER BY id;
    """
    work_sql = """
      SELECT w.id, w.work_title, w.site_id
      FROM works w
      WHERE w.deleted_at IS NULL AND w.site_id = 1152
      ORDER BY w.id;
    """
    emp_out = ssh_psql(emp_sql)
    work_out = ssh_psql(work_sql)

    employees = []
    for line in emp_out.splitlines():
        if not line.strip():
            continue
        parts = line.split("|")
        if len(parts) < 2:
            continue
        employees.append({
            "id": int(parts[0]),
            "fio": parts[1],
            "fio_key": norm_fio_key(parts[1]),
            "phone": parts[2] if len(parts) > 2 else "",
            "phone_norm": norm_phone(parts[2] if len(parts) > 2 else ""),
        })

    works = []
    for line in work_out.splitlines():
        if not line.strip():
            continue
        parts = line.split("|")
        if len(parts) < 2:
            continue
        works.append({
            "id": int(parts[0]),
            "work_title": parts[1],
            "site_id": int(parts[2]) if len(parts) > 2 and parts[2] else None,
        })
    return employees, works


def match_employee(fio_key: str, phone_norm: str, employees: list[dict]) -> dict | None:
    if phone_norm:
        by_phone = [e for e in employees if e.get("phone_norm") == phone_norm]
        if len(by_phone) == 1:
            return by_phone[0]
    exact = [e for e in employees if e["fio_key"] == fio_key]
    if len(exact) == 1:
        return exact[0]
    if len(exact) > 1:
        return exact[0]
    fam = fio_key.split("|")[0] if "|" in fio_key else fio_key
    partial = [e for e in employees if e["fio_key"].startswith(fam + "|")]
    if len(partial) == 1:
        return partial[0]
    return None


def should_skip_project(project_raw: str) -> bool:
    low = norm_fio(project_raw)
    if not low.strip():
        return True
    return any(p in low for p in SKIP_PROJECT_PATTERNS)


def match_work(project_raw: str, works: list[dict]) -> dict | None:
    kw = project_keyword(project_raw)
    if not kw:
        return None
    work_id = WORK_IDS.get(kw)
    if work_id:
        hit = next((w for w in works if w["id"] == work_id), None)
        if hit:
            return hit
    low_kw = kw.lower()
    matched = [w for w in works if low_kw in (w["work_title"] or "").lower()]
    if len(matched) == 1:
        return matched[0]
    if len(matched) > 1:
        main = [w for w in matched if "доп" not in (w["work_title"] or "").lower()]
        return main[0] if main else matched[0]
    return None


def build_plan_rows(personnel_path: Path, schedule_path: Path) -> list[dict]:
    personnel = read_personnel_xlsx(personnel_path)
    schedule = read_schedule_xlsx(schedule_path)
    employees, works = load_db_snapshot()

    plans = []
    for row in personnel:
        if should_skip_project(row["project_raw"]):
            plans.append({**row, "status": "skipped", "notes": ["не проект останова"]})
            continue
        emp = match_employee(row["fio_key"], row.get("phone_norm", ""), employees)
        work = match_work(row["project_raw"], works)
        sched = schedule.get(row["fio_key"], {})
        planned_from = row.get("planned_from_file")
        if planned_from:
            planned_from = planned_from.isoformat()
        elif sched.get("planned_from"):
            planned_from = sched["planned_from"]
        planned_to = sched.get("planned_to")

        status = "ok"
        notes = []
        if not emp:
            status = "employee_not_found"
        if not work:
            status = "work_not_found" if status == "ok" else status + "+work_not_found"
        if emp and work:
            dup = [p for p in plans if p.get("employee_id") == emp["id"]]
            if dup:
                status = "duplicate_in_file"

        plans.append({
            **row,
            "employee_id": emp["id"] if emp else None,
            "employee_fio_db": emp["fio"] if emp else None,
            "work_id": work["id"] if work else None,
            "work_title_db": work["work_title"] if work else None,
            "planned_from": planned_from,
            "planned_to": planned_to,
            "schedule_ranges": sched.get("ranges"),
            "status": status,
            "notes": notes,
        })
    return plans


def json_default(obj):
    if isinstance(obj, (date, datetime)):
        return obj.isoformat()
    raise TypeError(type(obj))


def apply_plans(plans: list[dict]) -> int:
    applied = 0
    key = paramiko.Ed25519Key.from_private_key_file(str(SSH_KEY))
    client = paramiko.SSHClient()
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    client.connect(HOST, username="root", pkey=key, timeout=30)

    for p in plans:
        if p["status"] != "ok":
            continue
        sql = f"""
        BEGIN;
        UPDATE employee_planned_engagements
          SET status='cancelled', cancelled_at=now(), cancelled_by={CREATED_BY}, updated_at=now()
          WHERE employee_id={p['employee_id']} AND status='active';
        INSERT INTO employee_planned_engagements
          (employee_id, work_id, planned_from, planned_to, note, status, created_by)
        VALUES (
          {p['employee_id']},
          {p['work_id']},
          {f"'{p['planned_from']}'" if p.get('planned_from') else 'NULL'},
          {f"'{p['planned_to']}'" if p.get('planned_to') else 'NULL'},
          'Импорт Excel {datetime.now().date().isoformat()}',
          'active',
          {CREATED_BY}
        );
        INSERT INTO worker_readiness_log
          (employee_id, old_status, new_status, comment, source, changed_by)
        SELECT id, readiness_status, readiness_status,
               'planned_set: work #{p['work_id']} ({p['project_raw']})', 'hr', {CREATED_BY}
        FROM employees WHERE id={p['employee_id']};
        COMMIT;
        """
        cmd = f'PGPASSWORD=123456789 psql -U asgard -d asgard_crm -v ON_ERROR_STOP=1 -c "{sql.strip()}"'
        _, stdout, stderr = client.exec_command(cmd, timeout=60)
        err = stderr.read().decode("utf-8", errors="replace")
        if err and "ERROR" in err.upper():
            p["status"] = "apply_error"
            p["notes"].append(err[:200])
        else:
            applied += 1

    client.close()
    return applied


def print_report(plans: list[dict]) -> None:
    ok = [p for p in plans if p["status"] == "ok"]
    bad = [p for p in plans if p["status"] != "ok"]
    print(f"\n=== Итого строк: {len(plans)} | готово к импорту: {len(ok)} | проблемы: {len(bad)} ===\n")
    print("--- Совпадения ---")
    for p in ok[:30]:
        print(
            f"  {p['fio']:28} → emp#{p['employee_id']} | {p['project_raw']:22} → work#{p['work_id']} "
            f"| {p['planned_from'] or '—'} … {p['planned_to'] or '—'}"
        )
    if len(ok) > 30:
        print(f"  … и ещё {len(ok) - 30}")
    if bad:
        print("\n--- Не сопоставлено / дубли ---")
        for p in bad:
            print(
                f"  row {p['row']:3} {p['fio']:28} | {p['project_raw']:22} | {p['status']} "
                f"| emp={p.get('employee_id')} work={p.get('work_id')}"
            )


def main():
    parser = argparse.ArgumentParser(description="Импорт планируемого привлечения из Excel")
    parser.add_argument("--personnel", type=Path, default=DEFAULT_PERSONNEL)
    parser.add_argument("--schedule", type=Path, default=DEFAULT_SCHEDULE)
    parser.add_argument("--dry-run", action="store_true", help="Только отчёт, без записи в БД")
    parser.add_argument("--apply", action="store_true", help="Записать планы в prod БД")
    parser.add_argument("--json-out", type=Path, default=ROOT / "tools" / "planned_import_report.json")
    args = parser.parse_args()

    if not args.personnel.exists():
        print(f"Файл не найден: {args.personnel}")
        sys.exit(1)
    if not args.schedule.exists():
        print(f"Файл не найден: {args.schedule}")
        sys.exit(1)
    if not args.dry_run and not args.apply:
        args.dry_run = True

    print(f"Персонал: {args.personnel.name}")
    print(f"График:   {args.schedule.name}")

    plans = build_plan_rows(args.personnel, args.schedule)
    print_report(plans)

    args.json_out.write_text(
        json.dumps(plans, ensure_ascii=False, indent=2, default=json_default),
        encoding="utf-8",
    )
    print(f"\nОтчёт JSON: {args.json_out}")

    if args.apply:
        n = apply_plans(plans)
        print(f"\nПрименено записей: {n}")
        args.json_out.write_text(
        json.dumps(plans, ensure_ascii=False, indent=2, default=json_default),
        encoding="utf-8",
    )


if __name__ == "__main__":
    main()
