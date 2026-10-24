# -*- coding: utf-8 -*-
"""Нормализация сводного табеля:

1) Баллы: все отметки, кроме 6 / 7 / 12, заменяются на 11 и красятся зелёным.
   Значения 6 / 7 / 12 сохраняются вместе с их исходной заливкой.
2) Должности приводятся к 6 видам:
      ПТО
      Монтажник (слесарь/жестянщик)
      Мастер ответственный
      Мастер
      Сварщик
      Руководитель проекта
3) Пересобираются итоги строк, лист «Сводка», легенда и лист «Источники».
"""
import io
import re
import datetime as dt

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DST = r'C:\Users\Nikita-ASGARD\Downloads\Табель_огневой_подогреватель_фев-окт_2026_сводный.xlsx'
REPORT = r'C:\Users\Nikita-ASGARD\ASGARD-CRM\tools\_normalize_report.txt'

KEEP_VALUES = {6, 7, 12}
NEW_VALUE = 11
GREEN = 'FF92D050'
RED_OLD = 'FFFF0000'

START = dt.date(2026, 2, 1)
NDAYS = 273
FIRST_DATA_ROW = 6

# свотчи для особых значений (доминирующие цвета в исходнике)
VALUE_SWATCH = {6: 'FFCFFAFE', 7: 'FFFEE2E2', 12: 'FFFFF3CD'}

POSITIONS = ['ПТО', 'Монтажник (слесарь/жестянщик)', 'Мастер ответственный',
             'Мастер', 'Сварщик', 'Руководитель проекта']

# исходное значение -> единая должность
POS_MAP = {
    'Мастер, Ответственный за проведение работ': 'Мастер ответственный',
    'Мастер-прораб': 'Мастер',
    'Мастер': 'Мастер',
    'мастер': 'Мастер',
    'Руководитель ПТО и СК': 'ПТО',
    'Инженер ПТО': 'ПТО',
    'ПТО': 'ПТО',
    'Руководитель направления МЛСП «Приразломная» / Заведующий распределением рабочего персонала':
        'Руководитель проекта',
    'РП': 'Руководитель проекта',
    'Сварщик': 'Сварщик',
    'Сварщик 6 р': 'Сварщик',
    'сварщик': 'Сварщик',
    'Слесарь-монтажник': 'Монтажник (слесарь/жестянщик)',
    'Слесарь-монтажник, Изолировщик': 'Монтажник (слесарь/жестянщик)',
    'слесарь': 'Монтажник (слесарь/жестянщик)',
    'Изолировщик': 'Монтажник (слесарь/жестянщик)',
    'Жестянщик': 'Монтажник (слесарь/жестянщик)',
    'Монтажник ТТ': 'Монтажник (слесарь/жестянщик)',
    'Монтажник труб-дов': 'Монтажник (слесарь/жестянщик)',
    'слесарь-сантехник': 'Монтажник (слесарь/жестянщик)',
}

# один человек в источнике назван «Монтажник ТТ (бригадир)» — должность та же
POS_MAP_BY_FIO = {
    'Шарафудинов Андрей Александрович (бригадир)': 'Монтажник (слесарь/жестянщик)',
}


def frgb(c):
    f = c.fill
    if f is None or f.patternType is None:
        return None
    try:
        r = f.fgColor.rgb
    except Exception:
        return None
    return r.upper() if isinstance(r, str) else None


def num(v):
    if isinstance(v, (int, float)):
        return int(v) if float(v) == int(v) else v
    return v


buf = io.StringIO()
wb = openpyxl.load_workbook(DST)
sh = wb['Табель']

# ── найти границы
day_cols = [c for c in range(4, 4 + NDAYS)]
last_day_col = day_cols[-1]
COL_DAYS, COL_PTS, COL_RUB = last_day_col + 1, last_day_col + 2, last_day_col + 3

total_row = None
for r in range(FIRST_DATA_ROW, sh.max_row + 1):
    if str(sh.cell(row=r, column=2).value or '').strip() == 'ИТОГО':
        total_row = r
        break
if total_row is None:
    raise RuntimeError('не найдена строка ИТОГО')
last_data_row = total_row - 1

# ── 1. нормализация баллов
replaced = 0
kept = {6: 0, 7: 0, 12: 0}
green_fill = PatternFill('solid', fgColor=GREEN)
before_vals = {}
after_vals = {}
per_row = {}
for r in range(FIRST_DATA_ROW, last_data_row + 1):
    days = 0
    pts = 0
    months = {}
    for i, c in enumerate(day_cols):
        cell = sh.cell(row=r, column=c)
        v = cell.value
        if not isinstance(v, (int, float)):
            continue
        before_vals[(r, c)] = num(v)
        iv = num(v)
        if iv in KEEP_VALUES:
            kept[iv] += 1
            newv = iv
            # цвет оставляем исходный (если его нет — ставим свотч значения)
            if frgb(cell) is None:
                cell.fill = PatternFill('solid', fgColor=VALUE_SWATCH[iv])
        else:
            newv = NEW_VALUE
            replaced += 1
            cell.fill = green_fill
        cell.value = newv
        cell.font = Font(name='Calibri', size=8)
        cell.alignment = Alignment(horizontal='center', vertical='center')
        after_vals[(r, c)] = newv
        days += 1
        pts += newv
        d = START + dt.timedelta(days=i)
        mk = '%04d-%02d' % (d.year, d.month)
        m = months.setdefault(mk, [0, 0])
        m[0] += 1
        m[1] += newv
    per_row[r] = {'days': days, 'pts': pts, 'months': months}
    sh.cell(row=r, column=COL_DAYS).value = days
    sh.cell(row=r, column=COL_PTS).value = num(pts)
    sh.cell(row=r, column=COL_RUB).value = '=%s%d*500' % (get_column_letter(COL_PTS), r)

# итоговая строка
sh.cell(row=total_row, column=COL_DAYS).value = '=SUM(%s%d:%s%d)' % (
    get_column_letter(COL_DAYS), FIRST_DATA_ROW, get_column_letter(COL_DAYS), last_data_row)
sh.cell(row=total_row, column=COL_PTS).value = '=SUM(%s%d:%s%d)' % (
    get_column_letter(COL_PTS), FIRST_DATA_ROW, get_column_letter(COL_PTS), last_data_row)
sh.cell(row=total_row, column=COL_RUB).value = '=SUM(%s%d:%s%d)' % (
    get_column_letter(COL_RUB), FIRST_DATA_ROW, get_column_letter(COL_RUB), last_data_row)

buf.write('=== 1. БАЛЛЫ\n')
buf.write('заменено на %d: %d ячеек\n' % (NEW_VALUE, replaced))
buf.write('сохранено 6/7/12: %s\n' % kept)
buf.write('всего ячеек: %d\n' % len(after_vals))

# ── 2. должности
people = []
pos_before = {}
r = FIRST_DATA_ROW
while r <= last_data_row:
    fio = str(sh.cell(row=r, column=2).value or '').strip()
    old = str(sh.cell(row=r, column=3).value or '').strip()
    new = POS_MAP_BY_FIO.get(fio) or POS_MAP.get(old)
    if new is None:
        raise RuntimeError('нет правила для должности %r (%s)' % (old, fio))
    sh.cell(row=r, column=3).value = new
    pos_before[fio] = old
    people.append({'row': r, 'fio': fio, 'pos': new, 'old': old})
    r += 1

used = sorted({p['pos'] for p in people})
buf.write('\n=== 2. ДОЛЖНОСТИ\n')
buf.write('видов после нормализации: %d -> %s\n' % (len(used), used))
bad = [p for p in used if p not in POSITIONS]
if bad:
    raise RuntimeError('появились лишние виды: %s' % bad)
from collections import Counter
cnt = Counter(p['pos'] for p in people)
for p in POSITIONS:
    buf.write('  %-34s %2d\n' % (p, cnt.get(p, 0)))
buf.write('\nпереименования:\n')
moved = Counter((p['old'], p['pos']) for p in people if p['old'] != p['pos'])
for (old, new), n in sorted(moved.items()):
    buf.write('  %-62s -> %-34s %d\n' % (old[:62], new, n))

# ── 3. лист «Сводка»
sm = wb['Сводка']
POS_COL = 3
sm_people = {}
r = 4
while True:
    fio = sm.cell(row=r, column=2).value
    if not fio or str(fio).strip() == 'ИТОГО':
        break
    sm_people[str(fio).strip()] = r
    r += 1
sm_total = r

# шапка: D -> справка об источнике
sm.cell(row=3, column=4).value = 'Исходно (СРМ)'

mkeys = ['2026-%02d' % m for m in range(2, 11)]
col = 5
month_cols = {}
for mk in mkeys:
    month_cols[mk] = (col, col + 1)
    col += 2
TOT_DAYS, TOT_PTS, TOT_RUB = col, col + 1, col + 2

for p in people:
    rr = sm_people.get(p['fio'])
    if rr is None:
        raise RuntimeError('нет строки на «Сводке» для %r' % p['fio'])
    pp = per_row[p['row']]
    sm.cell(row=rr, column=POS_COL).value = p['pos']
    sm.cell(row=rr, column=4).value = pos_before[p['fio']]
    for mk in mkeys:
        dc, pc = month_cols[mk]
        m = pp['months'].get(mk)
        sm.cell(row=rr, column=dc).value = (m[0] if m else None)
        sm.cell(row=rr, column=pc).value = (num(m[1]) if m else None)
    sm.cell(row=rr, column=TOT_DAYS).value = pp['days']
    sm.cell(row=rr, column=TOT_PTS).value = num(pp['pts'])
    sm.cell(row=rr, column=TOT_RUB).value = '=%s%d*500' % (get_column_letter(TOT_PTS), rr)

for c in range(5, TOT_RUB + 1):
    sm.cell(row=sm_total, column=c).value = '=SUM(%s4:%s%d)' % (
        get_column_letter(c), get_column_letter(c), sm_total - 1)

buf.write('\n=== 3. СВОДКА\n')
buf.write('людей: %d, итоговая строка %d\n' % (len(sm_people), sm_total))
for mk, name in zip(mkeys, ['Фев', 'Мар', 'Апр', 'Май', 'Июн', 'Июл', 'Авг', 'Сен', 'Окт']):
    d = sum(per_row[p['row']]['months'].get(mk, [0, 0])[0] for p in people)
    pts = sum(per_row[p['row']]['months'].get(mk, [0, 0])[1] for p in people)
    buf.write('  %s дней=%4d баллов=%s\n' % (name, d, num(pts)))

# ── 4. легенда
lr = total_row + 2
# очистить прежнюю легенду
for rr in range(lr, sh.max_row + 1):
    for cc in (2, 3):
        sh.cell(row=rr, column=cc).value = None
        sh.cell(row=rr, column=cc).fill = PatternFill(fill_type=None)

sh.cell(row=lr, column=2, value='ЛЕГЕНДА').font = Font('Calibri', size=10, bold=True)
legend = [
    (GREEN, '11 баллов — все отметки, кроме 6 / 7 / 12'),
    (VALUE_SWATCH[6], '6 баллов — дорога / самолёт / вертолёт'),
    (VALUE_SWATCH[7], '7 баллов — полдня / медосмотр'),
    (VALUE_SWATCH[12], '12 баллов — корабль / ночной переезд'),
]
for k, (rgb, label) in enumerate(legend, 1):
    sw = sh.cell(row=lr + k, column=2)
    sw.fill = PatternFill('solid', fgColor=rgb)
    sw.border = Border(left=Side(style='thin'), right=Side(style='thin'),
                       top=Side(style='thin'), bottom=Side(style='thin'))
    c = sh.cell(row=lr + k, column=3, value=label)
    c.font = Font('Calibri', size=9)

# ── 5. лист «Источники» — дописать про нормализацию
sr = wb['Источники и сверка']
add = [
    ('', ''),
    ('Нормализация баллов', 'все отметки, кроме 6 / 7 / 12, заменены на 11 и выделены зелёным '
                            '#92D050. Значения 6 / 7 / 12 сохранены с исходной заливкой. '
                            'Заменено %d ячеек; сохранено: 6 — %d, 7 — %d, 12 — %d. '
                            'Итого отметок: %d.'
        % (replaced, kept[6], kept[7], kept[12], len(after_vals))),
    ('Нормализация должностей', 'приведены к 6 видам: ' + '; '.join(POSITIONS) + '.'),
    ('Расчёт', '1 балл = 500 ₽. «Заработок» = Баллов × 500, пересчитан по новым значениям.'),
]
r = sr.max_row + 2
for a, b in add:
    ca = sr.cell(row=r, column=1, value=a)
    ca.font = Font('Calibri', size=9, bold=True)
    ca.alignment = Alignment(vertical='top')
    cb = sr.cell(row=r, column=2, value=b)
    cb.font = Font('Calibri', size=9)
    cb.alignment = Alignment(vertical='top', wrap_text=True)
    sr.row_dimensions[r].height = 46
    r += 1

wb.calculation.fullCalcOnLoad = True
wb.save(DST)

total_pts = sum(pp['pts'] for pp in per_row.values())
buf.write('\n=== ИТОГ\n')
buf.write('баллов всего: %s -> %s ₽\n' % (num(total_pts), num(total_pts * 500)))
buf.write('дней всего: %d\n' % sum(pp['days'] for pp in per_row.values()))
buf.write('файл: %s\n' % DST)

open(REPORT, 'w', encoding='utf-8').write(buf.getvalue())
print('saved')
