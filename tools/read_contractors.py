import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

import openpyxl
wb = openpyxl.load_workbook(r'C:\Users\Nikita-ASGARD\Downloads\export_contractors_18_06.xlsx', data_only=True)
for sh_name in wb.sheetnames:
    ws = wb[sh_name]
    print(f'=== Лист: {sh_name} ({ws.max_row} строк x {ws.max_column} колонок) ===')
    headers = [str(c.value) if c.value is not None else '' for c in ws[1]]
    print('Колонки:', headers)
    print()
    for i, row in enumerate(ws.iter_rows(min_row=2, max_row=8, values_only=True), 1):
        print(f'  строка {i+1}:', row)
    print()
