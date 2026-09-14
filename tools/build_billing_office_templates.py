# -*- coding: utf-8 -*-
"""Собрать docxtemplater-шаблоны счёта и акта по бланкам пользователя + логотип."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "templates" / "billing"
NAVY = RGBColor(0x1E, 0x4D, 0x8C)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)


def set_run(run, size=10, bold=False, color=BLACK, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def p_text(p, text, size=10, bold=False, color=BLACK, align=None, italic=False):
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, color=color, italic=italic)
    return run


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        if edge in kwargs:
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), kwargs[edge].get("val", "single"))
            el.set(qn("w:sz"), str(kwargs[edge].get("sz", 4)))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), kwargs[edge].get("color", "1E4D8C"))
            tcBorders.append(el)
    tcPr.append(tcBorders)


def shade(cell, fill="F3F4F6"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def thin():
    return {"val": "single", "sz": 4, "color": "1E4D8C"}


def cell_p(cell, text="", size=9, bold=False, color=BLACK, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    if align is not None:
        p.alignment = align
    if text:
        p_text(p, text, size=size, bold=bold, color=color)
    return p


def page_setup(doc):
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(1.5)
    sec.right_margin = Cm(1.2)
    sec.top_margin = Cm(1.2)
    sec.bottom_margin = Cm(1.2)


def add_logo(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    p_text(p, "{%logo_img}", size=10)


def build_invoice():
    doc = Document()
    page_setup(doc)
    add_logo(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_text(p, "Счёт на оплату № {number} от {date}", size=16, bold=True, color=NAVY)

    bank = doc.add_table(rows=4, cols=4)
    bank.autofit = True
    labels = [
        ("Банк получателя", "{bank_name}", "БИК", "{bank_bik}"),
        ("", "{bank_city}", "Сч. №", "{bank_ks}"),
        ("ИНН {inn}", "КПП {kpp}", "Сч. №", "{bank_rs}"),
        ("Получатель", "{company_full}", "", ""),
    ]
    for i, row in enumerate(labels):
        for j, val in enumerate(row):
            cell_p(bank.cell(i, j), val, size=8, bold=(j % 2 == 0 and val and "{" not in val))
            set_cell_border(bank.cell(i, j), top=thin(), left=thin(), bottom=thin(), right=thin())
    bank.cell(0, 0).merge(bank.cell(1, 0))
    bank.cell(3, 1).merge(bank.cell(3, 3))

    doc.add_paragraph()
    p = doc.add_paragraph()
    p_text(p, "Поставщик: ", size=10, bold=True)
    p_text(p, "{executor_line}", size=10)
    p = doc.add_paragraph()
    p_text(p, "Покупатель: ", size=10, bold=True)
    p_text(p, "{customer_line}", size=10)
    p = doc.add_paragraph()
    p_text(p, "Основание: ", size=10, bold=True)
    p_text(p, "{basis}", size=10)

    items = doc.add_table(rows=2, cols=6)
    headers = ["№", "Наименование", "Кол-во", "Ед.", "Цена", "Сумма"]
    for j, h in enumerate(headers):
        cell_p(items.cell(0, j), h, size=9, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(items.cell(0, j), "1E4D8C")
        set_cell_border(items.cell(0, j), top=thin(), left=thin(), bottom=thin(), right=thin())
    loop = ["{#items}{n}", "{name}", "{qty}", "{unit}", "{price}", "{sum}{/items}"]
    for j, val in enumerate(loop):
        cell_p(items.cell(1, j), val, size=9)
        set_cell_border(items.cell(1, j), top=thin(), left=thin(), bottom=thin(), right=thin())

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_text(p, "Итого: {amount} ₽", size=10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_text(p, "НДС {vat_pct}%: {vat_amount} ₽", size=10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_text(p, "Всего к оплате: {total} ₽", size=12, bold=True, color=NAVY)

    p = doc.add_paragraph()
    p_text(p, "Всего наименований {items_count}, на сумму {total_words}", size=10, italic=True)
    p = doc.add_paragraph()
    p_text(p, "Оплатить до: {due_date}", size=10)

    sig = doc.add_table(rows=3, cols=3)
    cell_p(sig.cell(0, 0), "{director_title}", size=10, bold=True)
    cell_p(sig.cell(0, 1), "{%signature_img}", size=10)
    cell_p(sig.cell(0, 2), "{%stamp_img}", size=10)
    cell_p(sig.cell(1, 0), "_________________ / {director} /", size=10)
    cell_p(sig.cell(2, 0), "Главный бухгалтер", size=10, bold=True)
    cell_p(sig.cell(2, 1), "_________________ / {accountant} /", size=10)

    out = OUT / "invoice-tpl.docx"
    doc.save(out)
    print("wrote", out)


def build_act():
    doc = Document()
    page_setup(doc)
    add_logo(doc)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_text(p, "АКТ", size=16, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_text(p, "сдачи-приёмки выполненных работ", size=14, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_text(p, "№ {number} от {date}", size=12, bold=True)

    p = doc.add_paragraph()
    p_text(p, "{customer_name}", size=10, bold=True)
    p_text(p, ", именуемое в дальнейшем «Заказчик», в лице {customer_signer}, с одной стороны, и ", size=10)
    p_text(p, "{company_full}", size=10, bold=True)
    p_text(
        p,
        ", именуемое в дальнейшем «Исполнитель», в лице {director_title} {director}, с другой стороны, составили настоящий акт о нижеследующем.",
        size=10,
    )

    p = doc.add_paragraph()
    p_text(p, "Исполнителем выполнены, а Заказчиком приняты следующие работы. Основание: {basis}", size=10)

    items = doc.add_table(rows=2, cols=6)
    headers = ["№ п/п", "Наименование выполненных работ", "Ед. изм.", "Кол-во", "Цена, руб.", "Сумма, руб."]
    for j, h in enumerate(headers):
        cell_p(items.cell(0, j), h, size=8, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), align=WD_ALIGN_PARAGRAPH.CENTER)
        shade(items.cell(0, j), "1E4D8C")
        set_cell_border(items.cell(0, j), top=thin(), left=thin(), bottom=thin(), right=thin())
    loop = ["{#items}{n}", "{name}", "{unit}", "{qty}", "{price}", "{sum}{/items}"]
    for j, val in enumerate(loop):
        cell_p(items.cell(1, j), val, size=9)
        set_cell_border(items.cell(1, j), top=thin(), left=thin(), bottom=thin(), right=thin())

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_text(p, "Итого: {amount} ₽", size=10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_text(p, "НДС {vat_pct}%: {vat_amount} ₽", size=10)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_text(p, "Всего выполнено работ на сумму: {total} ₽", size=12, bold=True, color=NAVY)
    p = doc.add_paragraph()
    p_text(p, "({total_words})", size=10, italic=True)

    p = doc.add_paragraph()
    p_text(
        p,
        "Качество выполненных работ проверено Заказчиком и соответствует условиям договора / заказа. Настоящий акт составлен в двух экземплярах, имеющих одинаковую юридическую силу, по одному для каждой из Сторон.",
        size=10,
    )

    p = doc.add_paragraph()
    p_text(p, "Подписи Сторон:", size=11, bold=True)

    sig = doc.add_table(rows=4, cols=2)
    cell_p(sig.cell(0, 0), "Работы сдал (Исполнитель)", size=10, bold=True)
    cell_p(sig.cell(0, 1), "Работы принял (Заказчик)", size=10, bold=True)
    cell_p(sig.cell(1, 0), "{%signature_img}", size=10)
    cell_p(sig.cell(1, 1), "", size=10)
    cell_p(sig.cell(2, 0), "{%stamp_img}", size=10)
    cell_p(sig.cell(2, 1), "М.П.", size=10)
    cell_p(sig.cell(3, 0), "_________________ / {director} /", size=10)
    cell_p(sig.cell(3, 1), "_________________ / {customer_signer} /", size=10)

    out = OUT / "act-tpl.docx"
    doc.save(out)
    print("wrote", out)


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    build_invoice()
    build_act()
