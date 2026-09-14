# -*- coding: utf-8 -*-
"""Build July+August CRM timesheet Excel from dumped CSVs."""
from __future__ import annotations

import csv
from collections import defaultdict
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(r"C:\Users\Nikita-ASGARD\ASGARD-CRM\tools")
OUT = Path(r"C:\Users\Nikita-ASGARD\Downloads\табель_CRM_июль_август_2026.xlsx")
POINT_VALUE = 500

C = {
    "road": "7030A0",
    "ship": "00B0F0",
    "op": "D60093",
    "w403": "2874A6",
    "w404": "E67E22",
    "w405": "1E8449",
    "w406": "148F77",
    "w407": "C0392B",
    "w418": "9A7D0A",
    "w353": "7F7F7F",
    "header": "1F4E79",
    "med": "F8CBAD",
    "wh": "C6E0B4",
    "remote": "FFC000",
}
WORK_HEX = {
    354: C["op"], 403: C["w403"], 404: C["w404"], 405: C["w405"],
    406: C["w406"], 407: C["w407"], 418: C["w418"], 353: C["w353"],
}
STAGE_LABEL = {
    "travel": "дор", "ship": "суд", "helicopter": "вер",
    "medical": "мед", "warehouse": "скл", "training": "уч",
    "waiting": "ож", "remote": "уд", "office": "оф",
}
STAGE_COLOR = {
    "travel": C["road"], "ship": C["ship"], "helicopter": C["road"],
    "medical": C["med"], "warehouse": C["wh"], "training": C["med"],
    "waiting": C["remote"], "remote": C["remote"], "office": "70AD47",
}
FONT = Font(name="Arial", size=8)
FONT_W = Font(name="Arial", size=8, bold=True, color="FFFFFF")
FONT_B = Font(name="Arial", size=9, bold=True, color="FFFFFF")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)
CENTER = Alignment(horizontal="center", vertical="center")
DARK = set(WORK_HEX.values()) | {C["road"], C["ship"], C["header"], "2E75B6", "C0007A"}


def fill(h):
    return PatternFill("solid", fgColor=h)


def daterange(a, b):
    d = a
    while d <= b:
        yield d
        d += timedelta(days=1)


def main():
    checkins = list(csv.DictReader((ROOT / "_crm_export_checkins.csv").open(encoding="utf-8")))
    stages = list(csv.DictReader((ROOT / "_crm_export_stages.csv").open(encoding="utf-8")))

    people = {}
    cells = defaultdict(dict)  # eid -> date -> (label, color)

    for r in checkins:
        eid = int(r["employee_id"])
        people[eid] = r["fio"]
        d = r["d"][:10]
        wid = int(r["work_id"]) if r.get("work_id") else None
        amt = float(r["amount_earned"] or 0)
        pts = int(round(amt / POINT_VALUE)) if amt else ""
        hx = WORK_HEX.get(wid, "B0B0B0")
        cells[eid][d] = (pts if pts != 0 else "Д", hx)

    for r in stages:
        eid = int(r["employee_id"])
        people[eid] = r["fio"]
        st = r["stage_type"]
        label = STAGE_LABEL.get(st, (st or "?")[:3])
        hx = STAGE_COLOR.get(st, C["road"])
        d0 = date.fromisoformat(r["d_from"][:10])
        d1 = date.fromisoformat((r["d_to"] or r["d_from"])[:10])
        for d in daterange(d0, d1):
            if d.year != 2026 or d.month not in (7, 8):
                continue
            cells[eid][d.isoformat()] = (label, hx)

    jul = [date(2026, 7, d) for d in range(1, 32)]
    aug = [date(2026, 8, d) for d in range(1, 32)]
    rows = sorted(people.items(), key=lambda x: x[1].lower())

    wb = Workbook()
    lg = wb.active
    lg.title = "Легенда"
    lg["A1"] = "Табель CRM июль + август 2026 — выгрузка после внесения табеля Хосе"
    lg["A1"].font = Font(name="Arial", size=12, bold=True, color="1F4E79")
    lg.merge_cells("A1:F1")
    lg["A2"] = (
        "Все сотрудники, у кого в CRM есть смена или этап за июль–август. "
        "Число в ячейке — баллы (сумма / 500). Цвет — объект. "
        "дор/суд/вер — этапы. Иванейкин в этот импорт не вносился."
    )
    lg["A2"].font = Font(name="Arial", size=9, italic=True, color="666666")
    lg.merge_cells("A2:F2")
    r = 4
    for hx, text in [
        (C["op"], "354 ОП"), (C["w403"], "403 Деаэратор"), (C["w404"], "404 Зачистка"),
        (C["w405"], "405 Каусорб"), (C["w406"], "406 Оголовок"), (C["w407"], "407 Ёмкостный парк"),
        (C["w418"], "418 сводная"), (C["road"], "дорога / вертолёт"), (C["ship"], "судно"),
    ]:
        lg.cell(r, 1).fill = fill(hx)
        lg.cell(r, 1).border = THIN
        lg.cell(r, 2, value=text).font = Font(name="Arial", size=9)
        r += 1
    lg.column_dimensions["A"].width = 14
    lg.column_dimensions["B"].width = 40

    ts = wb.create_sheet("Табель", 1)
    ts.cell(1, 1, value="№").font = FONT_B
    ts.cell(1, 1).fill = fill(C["header"])
    ts.cell(1, 2, value="ФИО").font = FONT_B
    ts.cell(1, 2).fill = fill(C["header"])
    ts.merge_cells(start_row=1, start_column=3, end_row=1, end_column=33)
    c = ts.cell(1, 3, value="Июль 2026")
    c.font = FONT_B
    c.fill = fill("2E75B6")
    c.alignment = CENTER
    ts.merge_cells(start_row=1, start_column=34, end_row=1, end_column=64)
    c = ts.cell(1, 34, value="Август 2026")
    c.font = FONT_B
    c.fill = fill("C0007A")
    c.alignment = CENTER
    ts.cell(2, 1).fill = fill(C["header"])
    ts.cell(2, 2).fill = fill(C["header"])
    for i, d in enumerate(jul, 3):
        cell = ts.cell(2, i, value=d.day)
        cell.font = FONT_B
        cell.fill = fill("2E75B6")
        cell.alignment = CENTER
        cell.border = THIN
        ts.cell(1, i).fill = fill("2E75B6")
        ts.cell(1, i).border = THIN
    for i, d in enumerate(aug, 34):
        cell = ts.cell(2, i, value=d.day)
        cell.font = FONT_B
        cell.fill = fill("C0007A")
        cell.alignment = CENTER
        cell.border = THIN
        ts.cell(1, i).fill = fill("C0007A")
        ts.cell(1, i).border = THIN

    for idx, (eid, fio) in enumerate(rows, 3):
        ts.cell(idx, 1, value=idx - 2).font = FONT
        ts.cell(idx, 1).border = THIN
        ts.cell(idx, 2, value=fio).font = FONT
        ts.cell(idx, 2).border = THIN
        for i, d in enumerate(jul + aug, 3):
            cell = ts.cell(idx, i)
            cell.alignment = CENTER
            cell.border = THIN
            cell.font = FONT
            rec = cells[eid].get(d.isoformat())
            if rec:
                label, hx = rec
                cell.value = label
                cell.fill = fill(hx)
                if hx in DARK:
                    cell.font = FONT_W
        ts.row_dimensions[idx].height = 16

    ts.column_dimensions["A"].width = 4
    ts.column_dimensions["B"].width = 32
    for c in range(3, 65):
        ts.column_dimensions[get_column_letter(c)].width = 3.6
    ts.freeze_panes = "C3"
    ts.auto_filter.ref = f"A2:BL{2 + len(rows)}"
    ts.sheet_view.showGridLines = False

    sm = wb.create_sheet("Сводка")
    sm["A1"] = f"Строк: {len(rows)} · смен: {len(checkins)} · этапов: {len(stages)}"
    sm["A1"].font = Font(name="Arial", size=11, bold=True)

    wb.save(OUT)
    print("saved", OUT, "people", len(rows), "checkins", len(checkins), "stages", len(stages))


if __name__ == "__main__":
    main()
