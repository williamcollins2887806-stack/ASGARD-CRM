# -*- coding: utf-8 -*-
"""Проверка сводного табеля: независимая пересборка эталона из источников."""
import io
import json
import re
import datetime as dt
from collections import Counter

import openpyxl

MINE = r'C:\Users\Nikita-ASGARD\Downloads\Табель_огневой_подогреватель_с_фев2026.xlsx'
CRM = r'C:\Users\Nikita-ASGARD\Downloads\timesheet_work354_2026-02-01_2026-10-31.xlsx'
MAP = r'C:\Users\Nikita-ASGARD\ASGARD-CRM\tools\_map_final.json'
DST = r'C:\Users\Nikita-ASGARD\Downloads\Табель_огневой_подогреватель_фев-окт_2026_сводный.xlsx'
OUT = r'C:\Users\Nikita-ASGARD\ASGARD-CRM\tools\_merged_verify.txt'

START = dt.date(2026, 2, 1)
NDAYS = (dt.date(2026, 10, 31) - START).days + 1
buf = io.StringIO()
fails = []
checks = []


def ok(name, cond, detail=''):
    checks.append((name, 'PASS' if cond else 'FAIL', detail))
    if not cond:
        fails.append(name)
    return cond


def frgb(c):
    f = c.fill
    if f is None or f.patternType is None:
        return None
    try:
        r = f.fgColor.rgb
    except Exception:
        return None
    return r.upper() if isinstance(r, str) else None


def key(s):
    s = str(s or '').replace('\n', ' ').replace('\u00a0', ' ')
    s = re.sub(r'\(.*?\)', ' ', s)
    s = re.sub(r'[.,]', ' ', s)
    s = re.sub(r'\s+', ' ', s).strip().lower().replace('ё', 'е')
    return s


mapping = json.load(open(MAP, encoding='utf-8'))
MK = {key(k): v for k, v in mapping.items()}
# в своде ФИО берутся канонические (из БД), а в маппинге ключ — написание
# источника. Обратный индекс по каноническому имени.
CANON = {}
for v in mapping.values():
    if v.get('id') is not None:
        CANON[key(v['fio'])] = v['id']

# ─────────── эталон: независимо читаем оба источника
exp = {}   # (emp_id, date) -> value
wb = openpyxl.load_workbook(MINE, data_only=False)
for sheet, mon in {'февр 2026': 2, 'март 2026': 3, 'апрель 2026': 4,
                   'май 2026': 5, 'июнь 2026': 6}.items():
    ws = wb[sheet]
    legend = max(r for r in range(1, ws.max_row + 1)
                 if isinstance(ws.cell(row=r, column=1).value, str)
                 and 'подогрев' in ws.cell(row=r, column=1).value.lower())
    for r in range(2, legend):
        fio = ws.cell(row=r, column=1).value
        if not fio:
            continue
        reds = [c for c in range(2, 33)
                if frgb(ws.cell(row=r, column=c)) == 'FFFF0000'
                and isinstance(ws.cell(row=r, column=c).value, (int, float))]
        if not reds:
            continue
        eid = MK[key(fio)]['id']
        for c in reds:
            exp[(eid, dt.date(2026, mon, int(ws.cell(row=1, column=c).value)))] = \
                ws.cell(row=r, column=c).value

wbc = openpyxl.load_workbook(CRM, data_only=True)
wsc = wbc['Табель']
for r in range(5, 39):
    fio = wsc.cell(row=r, column=2).value
    if not fio:
        continue
    eid = MK[key(fio)]['id']
    for i in range(NDAYS):
        cell = wsc.cell(row=r, column=3 + i)
        if isinstance(cell.value, (int, float)):
            k = (eid, START + dt.timedelta(days=i))
            if k in exp:
                continue
            exp[k] = cell.value

buf.write('эталон собрано независимо: %d отметок\n' % len(exp))

# ─────────── читаем результат
res = openpyxl.load_workbook(DST, data_only=False)
buf.write('листы: %s\n\n' % res.sheetnames)
sh = res['Табель']

# 1) структура
ok('лист «Табель» есть', 'Табель' in res.sheetnames)
ok('лист «Сводка» есть', 'Сводка' in res.sheetnames)

# найти строку данных
hdr_day_row = None
for r in range(1, 12):
    if str(sh.cell(row=r, column=4).value or '').startswith('01'):
        hdr_day_row = r
        break
ok('шапка дней найдена', hdr_day_row is not None, 'строка %s' % hdr_day_row)

first_data = hdr_day_row + 1
# заголовки дней
day_headers = []
for i in range(NDAYS):
    v = str(sh.cell(row=hdr_day_row, column=4 + i).value or '')
    day_headers.append(v.split('\n')[0])
expected = ['%02d' % (START + dt.timedelta(days=i)).day for i in range(NDAYS)]
ok('календарь дней 01.02..31.10', day_headers == expected,
   'первые: %s последние: %s' % (day_headers[:3], day_headers[-3:]))

# где итоги: заголовки объединены по строкам шапки, ищем по всей шапке
tot_cols = {}
for hdr_row in (hdr_day_row - 2, hdr_day_row - 1, hdr_day_row):
    for c in range(4 + NDAYS, 4 + NDAYS + 5):
        v = sh.cell(row=hdr_row, column=c).value
        if v and str(v).strip() in ('Дней', 'Баллов', 'Заработок, ₽'):
            tot_cols[str(v).strip()] = c
ok('есть столбцы Дней/Баллов/Заработок',
   all(k in tot_cols for k in ('Дней', 'Баллов')),
   str(tot_cols))

# 2) собрать фактические значения
got = {}
people_rows = {}
r = first_data
while True:
    fio = sh.cell(row=r, column=2).value
    if not fio or str(fio).strip() == 'ИТОГО':
        break
    pos = sh.cell(row=r, column=3).value
    people_rows[r] = (str(fio), pos)
    for i in range(NDAYS):
        cell = sh.cell(row=r, column=4 + i)
        if cell.value is None:
            continue
        if not isinstance(cell.value, (int, float)):
            fails.append('r%d день %d: нечисло %r' % (r, i + 1, cell.value))
            continue
        got[(r, i)] = cell.value
    r += 1
last_data = r - 1
total_row = r
buf.write('строк данных: %d (с %d по %d), итоговая строка %d\n'
          % (len(people_rows), first_data, last_data, total_row))

# 3) сверить по людям: сопоставляем ФИО -> emp_id
got_by_emp = {}
unmatched = []
for row, (fio, pos) in people_rows.items():
    info = MK.get(key(fio))
    eid = info['id'] if info else CANON.get(key(fio))
    if not eid:
        unmatched.append(fio)
        continue
    for i in range(NDAYS):
        if (row, i) in got:
            got_by_emp[(eid, START + dt.timedelta(days=i))] = got[(row, i)]
ok('все ФИО из свода опознаны', not unmatched, str(unmatched[:5]))

ok('число отметок совпадает с эталоном', len(got_by_emp) == len(exp),
   'свод=%d эталон=%d' % (len(got_by_emp), len(exp)))

missing = [k for k in exp if k not in got_by_emp]
extra = [k for k in got_by_emp if k not in exp]
valdiff = [(k, exp[k], got_by_emp[k]) for k in exp if k in got_by_emp and exp[k] != got_by_emp[k]]
ok('нет потерянных отметок', not missing, '%d %s' % (len(missing), missing[:5]))
ok('нет лишних отметок', not extra, '%d %s' % (len(extra), extra[:5]))
ok('значения совпадают', not valdiff, '%d %s' % (len(valdiff), valdiff[:5]))

# 4) суммы по строке
bad_tot = []
for row in people_rows:
    nums = [sh.cell(row=row, column=4 + i).value for i in range(NDAYS)]
    nums = [v for v in nums if isinstance(v, (int, float))]
    d_cell = sh.cell(row=row, column=tot_cols['Дней']).value
    p_cell = sh.cell(row=row, column=tot_cols['Баллов']).value
    if d_cell != len(nums):
        bad_tot.append((row, 'дней', d_cell, len(nums)))
    s = sum(nums)
    if p_cell is None or abs(float(p_cell) - s) > 1e-6:
        bad_tot.append((row, 'баллов', p_cell, s))
ok('итоги строк сходятся', not bad_tot, str(bad_tot[:4]))

# 5) должности заполнены
nopos = [(row, fio) for row, (fio, pos) in people_rows.items() if not pos]
ok('у всех заполнена должность', not nopos, '%d %s' % (len(nopos), nopos[:5]))

# 6) ИТОГО-строка
tot_days = sh.cell(row=total_row, column=tot_cols['Дней']).value
tot_pts = sh.cell(row=total_row, column=tot_cols['Баллов']).value
ok('итоговая строка есть', sh.cell(row=total_row, column=2).value == 'ИТОГО',
   repr(sh.cell(row=total_row, column=2).value))
real_days = sum(1 for _ in got)
buf.write('итоговая строка: дней=%r баллов=%r; фактически дней=%d баллов=%s\n'
          % (tot_days, tot_pts, real_days, sum(got.values())))

# 7) месяцы: покрытие по данным
buf.write('\n=== отметки по месяцам в своде\n')
mm = Counter()
for (eid, d) in got_by_emp:
    mm['%04d-%02d' % (d.year, d.month)] += 1
for k in sorted(mm):
    buf.write('  %s %d\n' % (k, mm[k]))

# 8) лист «Сводка»
sm = res['Сводка']
sm_people = 0
r = 4
while True:
    v = sm.cell(row=r, column=2).value
    if not v or str(v).strip() == 'ИТОГО':
        break
    sm_people += 1
    r += 1
ok('Сводка: все люди', sm_people == len(people_rows), '%d vs %d' % (sm_people, len(people_rows)))

# сумма баллов на Сводке
sm_tot_row = r
sm_pts_col = None
for c in range(1, sm.max_column + 1):
    if str(sm.cell(row=3, column=c).value or '').startswith('Баллов всего'):
        sm_pts_col = c
if sm_pts_col:
    sm_total = sm.cell(row=sm_tot_row, column=sm_pts_col).value
    buf.write('Сводка: итог баллов (формула) = %r, фактически %s\n' % (sm_total, sum(got.values())))

# 9) дублей людей в своде нет
fio_counts = Counter(key(f) for f, _ in people_rows.values())
dups = [f for f, c in fio_counts.items() if c > 1]
ok('нет дублей ФИО в своде', not dups, str(dups))

buf.write('\n=== ГЕЙТ\n')
for name, st, det in checks:
    buf.write('  [%s] %-46s %s\n' % (st, name, det[:110]))
buf.write('\nИТОГ: %s\n' % ('PASS' if not fails else 'FAIL: %s' % fails))

open(OUT, 'w', encoding='utf-8').write(buf.getvalue())
print('PASS' if not fails else 'FAIL')
