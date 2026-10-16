# -*- coding: utf-8 -*-
import sys, re, json
sys.stdout.reconfigure(encoding='utf-8')
from decimal import Decimal as D
import docx, fitz

DOCX = r"C:\Users\Nikita-ASGARD\Desktop\ТКП ЮНИКС Ильский НПЗ (Т-1301 + ПК-1301).docx"
PDF  = r"C:\Users\Nikita-ASGARD\Desktop\ТКП ЮНИКС Ильский НПЗ (Т-1301 + ПК-1301).pdf"

def norm(s):
    # normalize spaces incl nbsp, hyphens
    s = s.replace('\u00a0',' ').replace('\u2009',' ').replace('\u202f',' ')
    return s

d = docx.Document(DOCX)

# full text: paragraphs + tables + headers/footers + raw xml
paras = [norm(p.text) for p in d.paragraphs]
tbl_cells = []
for t in d.tables:
    for r in t.rows:
        for c in r.cells:
            tbl_cells.append(norm(c.text))
hf = []
for s in d.sections:
    for h in (s.header, s.first_page_header, s.even_page_header, s.footer, s.first_page_footer, s.even_page_footer):
        if h is None: continue
        for p in h.paragraphs: hf.append(norm(p.text))
        for t in h.tables:
            for r in t.rows:
                for c in r.cells: hf.append(norm(c.text))
docx_all = "\n".join(paras + tbl_cells + hf)

pdftxt = ""
doc = fitz.open(PDF)
for p in doc:
    pdftxt += norm(p.get_text("text")) + "\n"
pdf_all = pdftxt

# raw xml text for hidden/other parts
import zipfile
raw = ""
with zipfile.ZipFile(DOCX) as z:
    for n in z.namelist():
        if n.endswith('.xml'):
            raw += z.read(n).decode('utf-8', errors='replace')

def report(name, needles, hay=docx_all, mode="contains"):
    hits = []
    for n in needles:
        if mode=="regex":
            if re.search(n, hay): hits.append(n)
        else:
            if n in hay: hits.append(n)
    print(f"[{name}] hits={len(hits)} -> {hits}")
    return hits

print("=== CRIT 5 ARITHMETIC (Decimal) ===")
def mul(a,b): return (D(a)*D(b))
checks = [
 ("47212.44*64", D("47212.44")*64, D("3021596.16")),
 ("15546.08*40", D("15546.08")*40, D("621843.20")),
 ("52511.19*64", D("52511.19")*64, D("3360716.16")),
 ("75758.84*20", D("75758.84")*20, D("1515176.80")),
 ("SUM no VAT", D("3643439.36")+D("3360716.16")+D("1515176.80"), D("8519332.32")),
 ("VAT1 22%", (D("3643439.36")*D("0.22")).quantize(D("0.01")), D("801556.66")),
 ("VAT2 22%", (D("3360716.16")*D("0.22")).quantize(D("0.01")), D("739357.56")),
 ("VAT3 22%", (D("1515176.80")*D("0.22")).quantize(D("0.01")), D("333338.90")),
 ("SUM VAT", D("801556.66")+D("739357.56")+D("333338.90"), D("1874253.12")),
 ("TOTAL with VAT", D("8519332.32")+D("1874253.12"), D("10393585.44")),
 ("B1 with VAT", D("3643439.36")+D("801556.66"), D("4444996.02")),
 ("B2 with VAT", D("3360716.16")+D("739357.56"), D("4100073.72")),
 ("B3 with VAT", D("1515176.80")+D("333338.90"), D("1848515.70")),
]
for n,a,b in checks:
    print(f"  {n}: {a} == {b} -> {'OK' if a==b else 'FAIL'}")

print("\n=== CRIT 6 MTO/INTERNAL ===")
report("C6-docx", ["Материально-технич","МТО","1 623 899","381 000","экономи","дублирующ","приобретается однократно","не уменьшается","не пересчитывается","раздельн"])
report("C6-pdf",  ["Материально-технич","МТО","1 623 899","381 000","экономи","дублирующ","приобретается однократно","не уменьшается","не пересчитывается","раздельн"], pdf_all)

print("\n=== CRIT 7 LEAKS ===")
leaks = ["себестоимость","маржа","прибыль","ФОТ","налог","накладные","непредвиденные","пайковые","зарплата","склад","логистика","проезд","проживание","командировочные","Саратов","Москва","спецтехник",
 "15 000","10 000","7 500","1 812 000","3 754 328","2 666 666","160 000","1 000 000","150 000","170 000","586 451","1 545 101","1 051 400","55 %"]
for h,nm in ((docx_all,"docx"),(pdf_all,"pdf")):
    hits=[]
    for n in leaks:
        if n in h: hits.append(n)
    print(f"  {nm}: {hits}")
# оклад except прокладок
for nm,h in (("docx",docx_all),("pdf",pdf_all)):
    occ=[m.start() for m in re.finditer("оклад", h)]
    bad=[h[max(0,i-1):i+6] for i in occ if not h[max(0,i-3):i+len("оклад")].endswith("проклад")]
    print(f"  оклад[{nm}] contexts_bad={bad}")

print("\n=== CRIT 8 SIZES ===")
for nm,h in (("docx",docx_all),("pdf",pdf_all)):
    for pat in [r'2\s*1/4', r'1\s*5/8', r'1\s*1/2', r'(?<![1/0-9])5/8', r'(?<![\d/])1\s*"\s']:
        print(f"  {nm} /{pat}/ ->", re.findall(pat,h))

print("\n=== CRIT 9 TERMINOLOGY ===")
report("C9-docx", ["заболчивание","потяжка","ANSIB","электродвиг ател","слесаря- монтажник","атап","2025","17-23.11.2027"])
report("C9-pdf",  ["заболчивание","потяжка","ANSIB","электродвиг ател","слесаря- монтажник","атап","2025","17-23.11.2027"], pdf_all)
print("  U+FFFD docx:", docx_all.count("\ufffd"), "raw:", raw.count("\ufffd"), "pdf:", pdf_all.count("\ufffd"))
print("  'забалчивание' count docx:", docx_all.count("забалчивание"), "pdf:", pdf_all.count("забалчивание"))
print("  'протяжка' docx:", len(re.findall("протяж", docx_all)), "pdf:", len(re.findall("протяж", pdf_all)))
print("  'ANSI B18.2.1' docx:", docx_all.count("ANSI B18.2.1"))
print("  'электродвигател' docx:", docx_all.count("электродвигател"), "pdf:", pdf_all.count("электродвигател"))
print("  'слесаря-монтажника' docx:", docx_all.count("слесаря-монтажника"), "pdf:", pdf_all.count("слесаря-"))

print("\n=== CRIT 1 PROHODY ===")
if False: pass
report("C1-forbidden", [r'четыре проход', r'в пять проходов', r'\bпять проходов'])
for t in d.tables:
    for r in t.rows:
        for c in r.cells:
            if c.text.strip() in ("4","5"):
                pass
# find "Число проходов"/"Проходов" columns
for ti,t in enumerate(d.tables):
    hdr=[c.text.strip() for c in t.rows[0].cells]
    if any(x in hdr for x in ("Число проходов","Проходов")):
        idx=[i for i,x in enumerate(hdr) if x in ("Число проходов","Проходов")]
        for r in t.rows[1:]:
            cells=[c.text.strip() for c in r.cells]
            print(f"  table{ti} hdr={hdr} -> values={[cells[i] for i in idx if i<len(cells)]}")

print("\n=== CRIT 11 ANTI-AI ===")
anti = ["важно отметить","следует отметить","стоит отметить","таким образом","в заключение","в целом","в свою очередь","вместе с тем","тем не менее","подводя итог","как известно","в рамках","в целях","в части","надлежащ","вышеуказанн","нижеследующ","осуществля","является","настоящим","ключев","оптимальн","комплексн","индивидуальн","гибкий подход","клиентоориентированн","широкий спектр","играет важную роль","позволит","обеспечит","способствует","на сегодняшний день","на данный момент","TBD","TODO","lorem","XXX","не заполнено","уточняется","пример","в области"]
for nm,h in (("docx",docx_all),("pdf",pdf_all)):
    hits=[a for a in anti if a in h]
    print(f"  {nm}: {hits}")

print("\n=== CRIT 2 STAGES ===")
for nm,h in (("docx",docx_all),):
    for sect in ["Раздел 1","Раздел 2"]:
        pass
st1=[p for p in paras if re.match(r'Этап \d', p)]
print("  stages order:", [re.match(r'Этап \d',p).group() for p in st1])

print("\n=== CRIT 4 OFFER ===")
for p in paras:
    if "второго выезда" in p: print("  FOUND:", p)

print("\n=== TABLE WIDTHS (docx) ===")
from docx.oxml.ns import qn
for ti,t in enumerate(d.tables):
    # grid cols
    grid = t._tbl.find(qn('w:tblGrid'))
    cols=[]
    if grid is not None:
        for gc in grid.findall(qn('w:gridCol')):
            w=gc.get(qn('w:w'))
            cols.append(int(w) if w else 0)
    tw = None
    tblPr = t._tbl.find(qn('w:tblPr'))
    if tblPr is not None:
        twel = tblPr.find(qn('w:tblW'))
        if twel is not None:
            tw = (twel.get(qn('w:w')), twel.get(qn('w:type')))
    total_twips = sum(cols)
    print(f"  table{ti}: cols={len(cols)} sum_twips={total_twips} cm={total_twips/566.929:.2f} tblW={tw}")

print("\n=== PDF PAGE COUNT ===", doc.page_count, "| DOCX pages(PDF-rendered) ===")
