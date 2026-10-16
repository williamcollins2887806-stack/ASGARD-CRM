# -*- coding: utf-8 -*-
import sys, json, io
sys.stdout.reconfigure(encoding='utf-8')
import docx
from docx.document import Document as _Doc
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

PATH = r"C:\Users\Nikita-ASGARD\Desktop\ТКП ЮНИКС Ильский НПЗ (Т-1301 + ПК-1301).docx"
d = docx.Document(PATH)

def iter_block_items(parent):
    if isinstance(parent, _Doc):
        parent_elm = parent.element.body
    else:
        parent_elm = parent._element
    for child in parent_elm.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)

out = []
pi = 0
ti = 0
for block in iter_block_items(d):
    if isinstance(block, Paragraph):
        out.append({"k":"P","i":pi,"s":block.style.name if block.style else "","t":block.text})
        pi += 1
    else:
        rows = []
        for r in block.rows:
            rows.append([c.text for c in r.cells])
        out.append({"k":"T","i":ti,"rows":rows})
        ti += 1

with open(r"C:\Users\Nikita-ASGARD\ASGARD-CRM\tmp_verify_tkp\docx_body.json","w",encoding="utf-8") as f:
    json.dump(out,f,ensure_ascii=False,indent=1)

# sections: headers/footers
secs=[]
for si,s in enumerate(d.sections):
    hdr=[]
    for h in [s.header, s.first_page_header, s.even_page_header]:
        if h is None: continue
        for p in h.paragraphs:
            if p.text.strip(): hdr.append(p.text)
        for t in h.tables:
            for r in t.rows:
                hdr.append(" | ".join(c.text for c in r.cells))
    ftr=[]
    for h in [s.footer, s.first_page_footer, s.even_page_footer]:
        if h is None: continue
        for p in h.paragraphs:
            if p.text.strip(): ftr.append(p.text)
        for t in h.tables:
            for r in t.rows:
                ftr.append(" | ".join(c.text for c in r.cells))
    secs.append({"section":si,"header":hdr,"footer":ftr})
with open(r"C:\Users\Nikita-ASGARD\ASGARD-CRM\tmp_verify_tkp\docx_hf.json","w",encoding="utf-8") as f:
    json.dump(secs,f,ensure_ascii=False,indent=1)

print("paragraphs:",pi,"tables:",ti,"sections:",len(secs))
